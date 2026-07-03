from __future__ import annotations

import os
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt


def set_run_font(run, name: str = '宋体', size_pt: int = 12) -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(text, style=f'Heading {level}')
    for r in p.runs:
        set_run_font(r)


def add_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style='Normal')
    for r in p.runs:
        set_run_font(r)


def main() -> None:
    out_path = Path(os.environ.get('OUT_PATH', 'user_manual.docx')).resolve()

    software_cn = os.environ.get('SOFTWARE_CN', 'Wonder Payment For WooCommerce（WooCommerce 支付网关插件）')
    software_en = os.environ.get('SOFTWARE_EN', 'Wonder Payment For WooCommerce')
    version = os.environ.get('SOFTWARE_VERSION', '1.0.4')

    doc = Document()

    # Cover / Title (simple, like the sample)
    p = doc.add_paragraph(software_cn)
    p.style = 'Normal'
    for r in p.runs:
        set_run_font(r, size_pt=16)

    p = doc.add_paragraph('用户手册')
    p.style = 'Normal'
    for r in p.runs:
        set_run_font(r, size_pt=16)

    doc.add_paragraph('')

    # Table of contents placeholder (Word will update fields manually)
    add_paragraph(doc, '（目录可在 Word 中通过“引用 → 目录”自动生成/更新）')
    doc.add_page_break()

    # Body
    p = doc.add_paragraph('用户手册')
    for r in p.runs:
        set_run_font(r, size_pt=14)

    # 1 引言
    add_heading(doc, '1引言', 1)

    add_heading(doc, '1.1编写目的', 2)
    add_paragraph(
        doc,
        f'本手册用于说明 {software_en}（版本 {version}）的安装、配置与使用方法，帮助电商商户在 WooCommerce 中接入 Wonder Payments，完成收款、退款以及订单状态同步等操作。'
    )

    add_heading(doc, '1.2背景', 2)
    add_paragraph(
        doc,
        'WooCommerce 是 WordPress 生态中广泛使用的电商插件。为满足商户对多渠道在线收款的需求，本软件以 WooCommerce 支付网关插件形式提供 Wonder Payments 支付能力，包括支付链接创建、Webhook 订单状态回传、手动同步以及退款能力等。'
    )

    add_heading(doc, '1.3定义', 2)
    add_paragraph(doc, 'WooCommerce：WordPress 电商插件框架。')
    add_paragraph(doc, '支付网关：在结算时为订单提供支付方式并对接第三方支付平台的组件。')
    add_paragraph(doc, 'Webhook：第三方服务向商户站点回调通知订单/支付状态变化的机制。')

    add_heading(doc, '1.4参考资料', 2)
    add_paragraph(doc, 'WordPress 官方文档（WordPress 插件开发与发布规范）。')
    add_paragraph(doc, 'WooCommerce 开发者文档（支付网关接口与扩展开发）。')
    add_paragraph(doc, 'Wonder Payments 开放接口文档（用于创建订单、查询、退款等）。')

    # 2 用途
    add_heading(doc, '2用途', 1)

    add_heading(doc, '2.1功能', 2)
    add_paragraph(doc, '本软件的主要功能包括：')
    add_paragraph(doc, '（1）在 WooCommerce 结算页提供 Wonder Payments 支付方式。')
    add_paragraph(doc, '（2）支持商户通过 App ID 等凭据完成网关配置与连通性校验。')
    add_paragraph(doc, '（3）创建 Wonder 侧支付订单与支付链接，并将结果写入 WooCommerce 订单。')
    add_paragraph(doc, '（4）接收并验证 Wonder Webhook 回调，实现支付状态自动更新。')
    add_paragraph(doc, '（5）提供后台手动同步订单状态能力。')
    add_paragraph(doc, '（6）支持全额/部分退款，并同步 WooCommerce 订单状态。')

    add_heading(doc, '2.2性能', 2)

    add_heading(doc, '2.2.1精度', 3)
    add_paragraph(doc, '金额计算以 WooCommerce 订单金额与币种为准；与第三方接口交互过程中的金额精度遵循接口约定。')

    add_heading(doc, '2.2.2时间特性', 3)
    add_paragraph(doc, '订单状态变更依赖 Webhook 到达与站点处理速度；在网络正常情况下可实现秒级到分钟级更新。')

    add_heading(doc, '2.2.3灵活性', 3)
    add_paragraph(doc, '支持多环境（生产/测试）切换；支持启用/停用、标题与描述等网关展示配置。')

    add_heading(doc, '2.3安全保密', 2)
    add_paragraph(doc, '插件对 Webhook 请求进行必要的来源与签名验证；商户密钥等敏感信息按照 WordPress/WooCommerce 设置项存储与输入过滤要求处理。')

    # 3 运行环境
    add_heading(doc, '3运行环境', 1)

    add_heading(doc, '3.1硬设备', 2)
    add_paragraph(doc, '普通云服务器/虚拟主机即可满足运行需求；建议具备稳定的公网访问能力以接收 Webhook 回调。')

    add_heading(doc, '3.2支持软件', 2)
    add_paragraph(doc, '操作系统：Linux/Windows/macOS（服务器端一般为 Linux）。')
    add_paragraph(doc, '运行平台：WordPress 5.8 及以上。')
    add_paragraph(doc, '依赖组件：WooCommerce 插件（必需）。')
    add_paragraph(doc, 'PHP：7.4 及以上。')

    add_heading(doc, '3.3数据结构', 2)
    add_paragraph(doc, '订单数据以 WooCommerce 订单表与订单元数据为主；插件按需在订单元数据中记录 Wonder 侧订单号、交易 UUID、参考号等字段。')

    # 4 使用过程
    add_heading(doc, '4使用过程', 1)

    add_heading(doc, '4.1安装与初始化', 2)
    add_paragraph(doc, '（1）将插件上传至 WordPress 的 /wp-content/plugins/ 目录或通过后台“插件”上传安装。')
    add_paragraph(doc, '（2）在 WordPress 后台启用插件。')
    add_paragraph(doc, '（3）进入 WooCommerce → 设置 → 付款，找到 Wonder Payments 并点击“管理”。')

    add_heading(doc, '4.2输入', 2)

    add_heading(doc, '4.2.1输入数据的现实背景', 3)
    add_paragraph(doc, '商户需要提供 Wonder 侧的商户账号与应用凭据，用于在下单与退款时调用第三方接口。')

    add_heading(doc, '4.2.2输入格式', 3)
    add_paragraph(doc, '在 WooCommerce 支付网关设置页中输入 App ID、环境选项等配置项；部分配置可能以多行文本方式输入（如 PEM 格式密钥）。')

    add_heading(doc, '4.2.3输入举例', 3)
    add_paragraph(doc, 'App ID：示例格式为字符串（具体值以商户后台分配为准）。')

    add_heading(doc, '4.3输出对每项输出作出说明', 2)

    add_heading(doc, '4.3.1输出数据的现实背景', 3)
    add_paragraph(doc, '当用户提交订单并选择 Wonder Payments 支付后，系统会生成支付链接/订单号，并在 WooCommerce 订单中记录关联信息。')

    add_heading(doc, '4.3.2输出格式', 3)
    add_paragraph(doc, 'WooCommerce 订单状态：pending/on-hold/processing/completed/refunded/failed 等；同时可能包含 Wonder 订单号、交易 UUID 等扩展字段。')

    add_heading(doc, '4.3.3输出举例', 3)
    add_paragraph(doc, '示例：订单状态从 pending 更新为 processing（Webhook 回调确认支付完成）。')

    add_heading(doc, '4.4文卷查询', 2)
    add_paragraph(doc, '商户可在 WooCommerce 订单列表查看订单支付状态与订单备注；插件提供必要的状态同步入口以便查询最新第三方状态。')

    add_heading(doc, '4.5出错处理和恢复', 2)
    add_paragraph(doc, '当第三方接口不可用或签名校验失败时，插件会记录日志并向管理员提示；商户可在恢复后通过手动同步或重试完成状态更新。')

    add_heading(doc, '4.6终端操作', 2)
    add_paragraph(doc, '本软件为 WordPress 插件，一般无需终端操作。开发/排障时可查看站点日志与插件 debug.log（如已启用）。')

    add_paragraph(doc, '')
    add_paragraph(doc, f'文档生成日期：{date.today().isoformat()}')

    doc.save(str(out_path))


if __name__ == '__main__':
    main()
