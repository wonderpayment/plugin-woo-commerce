from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path

from docx import Document


@dataclass(frozen=True)
class FileLine:
    path: str
    line_no: int
    text: str


EXCLUDE_DIRS = {
    '.git',
    '.idea',
    'vendor',
    '_copyright_out',
}

INCLUDE_SUFFIXES = {
    # Scope B: PHP + JS/CSS (exclude docs/config files)
    '.php', '.js', '.css',
}

EXCLUDE_FILES = {
    # Exclude non-source materials for soft copyright code printouts.
    'README.md',
    'readme.txt',
    'composer.json',
    'docs/wporg-review-progress.md',
}

CORE_ONLY_FILES = {
    # Scope 1: core payment logic only (entry + gateway implementation)
    'wonderpay-gateway-for-woocommerce.php',
    'includes/class-wonder-payments-gateway.php',
}


def iter_source_files(root: Path) -> list[Path]:
    files: list[Path] = []
    for p in root.rglob('*'):
        if not p.is_file():
            continue
        rel = p.relative_to(root)
        if any(part in EXCLUDE_DIRS for part in rel.parts):
            continue
        if str(rel).replace("\\", "/") in EXCLUDE_FILES:
            continue
        rel_posix = str(rel).replace("\\", "/")
        scope = os.environ.get("CODE_SCOPE", "B").strip().upper()
        if scope == "CORE_ONLY" and rel_posix not in CORE_ONLY_FILES:
            continue

        if p.suffix.lower() not in INCLUDE_SUFFIXES:
            continue
        files.append(p)
    files.sort(key=lambda x: str(x.relative_to(root)).lower())
    return files


def read_lines(p: Path) -> list[str]:
    try:
        return p.read_text(encoding='utf-8').splitlines()
    except UnicodeDecodeError:
        return p.read_text(encoding='utf-8', errors='replace').splitlines()


def build_all_lines(root: Path) -> list[FileLine]:
    out: list[FileLine] = []
    for f in iter_source_files(root):
        rel = str(f.relative_to(root))
        for idx, line in enumerate(read_lines(f), start=1):
            out.append(FileLine(path=rel, line_no=idx, text=line))
    return out


def choose_front_back_lines(
    file_lines: list[FileLine],
    *,
    front_pages: int = 30,
    back_pages: int = 30,
) -> list[FileLine]:
    # The sample code template uses docGrid linePitch=312 twips.
    # With A4 and 1" top/bottom margins, it fits about 44 lines per page.
    lines_per_page = 44
    front_budget = front_pages * lines_per_page
    back_budget = back_pages * lines_per_page

    # Word will wrap long lines, which inflates real page count.
    # To better approximate "30 pages + 30 pages" when rendered, count an
    # approximate wrapped-line cost per source line.
    #
    # For the template (Times New Roman 10.5pt, A4, ~1.25\" side margins),
    # ~80 characters per physical line is a conservative approximation.
    chars_per_physical_line = int(os.environ.get("CHARS_PER_LINE", "80"))

    def physical_cost(text: str) -> int:
        if not text:
            return 1
        # Treat tabs as 4 spaces for wrapping approximation.
        normalized = text.replace("\t", "    ")
        return max(1, (len(normalized) + chars_per_physical_line - 1) // chars_per_physical_line)

    def take_prefix(lines: list[FileLine], budget: int) -> list[FileLine]:
        out: list[FileLine] = []
        used = 0
        for fl in lines:
            cost = physical_cost(fl.text)
            if used + cost > budget and out:
                break
            out.append(fl)
            used += cost
        return out

    def take_suffix(lines: list[FileLine], budget: int) -> list[FileLine]:
        out_rev: list[FileLine] = []
        used = 0
        for fl in reversed(lines):
            cost = physical_cost(fl.text)
            if used + cost > budget and out_rev:
                break
            out_rev.append(fl)
            used += cost
        return list(reversed(out_rev))

    # If small enough, keep all.
    total_cost = sum(physical_cost(fl.text) for fl in file_lines)
    if total_cost <= front_budget + back_budget:
        return file_lines

    return take_prefix(file_lines, front_budget) + take_suffix(file_lines, back_budget)


def clear_paragraph(paragraph) -> None:
    # remove all runs
    p_el = paragraph._element
    for r in list(p_el):
        p_el.remove(r)


def main() -> None:
    project_root = Path(os.environ.get('PROJECT_ROOT', '.')).resolve()
    template_path = Path(os.environ['TEMPLATE_PATH']).resolve()
    out_path = Path(os.environ['OUT_PATH']).resolve()
    pages_mode = os.environ.get('PAGES_MODE', 'front_back').strip().lower()

    lines = build_all_lines(project_root)
    if pages_mode in {'front_back', 'first_last'}:
        lines = choose_front_back_lines(lines)

    doc = Document(str(template_path))

    # Replace content while preserving Normal style defaults from template.
    needed = len(lines)
    current = len(doc.paragraphs)

    # Ensure we have enough paragraphs
    if current < needed:
        for _ in range(needed - current):
            doc.add_paragraph('', style='Normal')

    # Fill paragraphs
    for i, fl in enumerate(lines):
        p = doc.paragraphs[i]
        # keep paragraph style as-is (Normal)
        clear_paragraph(p)
        p.add_run(fl.text)

    # Remove extra paragraphs at end (keep exactly needed)
    # Deleting paragraphs in python-docx: remove element from body.
    if len(doc.paragraphs) > needed:
        body = doc._element.body
        # doc.paragraphs is dynamic; remove from end
        for p in list(doc.paragraphs)[needed:][::-1]:
            body.remove(p._element)

    doc.save(str(out_path))


if __name__ == '__main__':
    main()
